"""Discrete Summary build: transcription workbook + .btl files -> xlsx + README.

One summary row per (cast, Niskin, replicate index) for every Niskin on the
workbook's Niskin Log. CTD values come from the .btl bottle with the same
position; the start position comes from the .btl NMEA header and is checked
against the Cast Log. Lab results (asset_discrete.labs) are joined to the
drawn samples by cast + sample number; a received result gets discrete flag
bit 2, a drawn sample with no result is "drawn, not received" (bit 1), and
analyses not drawn from a bottle get bit 9.

Outputs, named <Expedition>_<Cruise>_Discrete_Summary.*:

- .xlsx -- formatted for review. Numbers are stored as numbers with a number
  format matching their text, flags and times as text, so Excel's "CSV UTF-8"
  save reproduces the release CSV exactly.
- _README.txt -- revision history, general notes, published summary notes
  and the flag tables, in the historic EA layout.
- _build_log.txt -- cross-check results (positions, times, calibration age).
"""
from __future__ import annotations

import csv
import math
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from asset_discrete import REFERENCE
from asset_discrete.btl import CTD_COLUMNS, SENSOR_NAMES, BtlFile, read_btl
from asset_discrete.flags import (FLAG_TYPES, MATCH_BIT, NOT_DRAWN_BIT, NOT_RECEIVED_BIT, NOTES_BIT, OK_BIT,
                                  OLD_CAL_BIT, bits_by_text, definitions, flag)
from asset_discrete.labs import LabResult, read_lab
from asset_discrete.workbook import read_workbook

FILL = '-9999999'
POSITION_TOLERANCE_KM = 1.0
TIME_TOLERANCE_MIN = 15
CAL_MAX_AGE_DAYS = 365
DRAW_TEMP_TOLERANCE_C = 0.2

# CTD parameter flag column -> (sensor key for the calibration-age check, analysis giving bit 7)
CTD_FLAGS = {'CTD Pressure Flag': ('P', None), 'CTD Temperature 1 Flag': ('T0', None),
             'CTD Temperature 2 Flag': ('T1', None), 'CTD Conductivity 1 Flag': ('C0', 'salinity'),
             'CTD Conductivity 2 Flag': ('C1', 'salinity'), 'CTD Oxygen Flag': ('O2', 'oxygen'),
             'CTD Fluorescence Flag': ('FL', 'chlorophyll'), 'CTD Transmissometer Flag': ('TR', None),
             'CTD pH Flag': ('PH', None)}


@dataclass(frozen=True)
class Discrete:
    """One discrete measurement group in the summary."""
    source: str | None      # Niskin Log column with its sample numbers (None: never drawn by EA)
    values: tuple[str, ...]  # summary value columns
    flag: str               # summary flag column
    replicate: str          # summary replicate flag column
    issue_name: str | None  # Analysis name used on the Sample Issues sheet


DISCRETE = {
    'oxygen': Discrete('Oxygen Samples', ('Discrete Oxygen [mL/L]',), 'Discrete Oxygen Flag',
                       'Discrete Oxygen Replicate Flag', 'oxygen'),
    'chlorophyll': Discrete('Chlorophyll Filters', ('Discrete Chlorophyll [ug/L]', 'Discrete Phaeopigment [ug/L]',
                                                    'Discrete Fo/Fa Ratio'),
                            'Discrete Fluorescence Flag', 'Discrete Fluorescence Replicate Flag', 'chlorophyll'),
    'nutrients': Discrete('Nutrient Samples', ('Discrete Phosphate [uM]', 'Discrete Silicate [uM]',
                                               'Discrete Nitrate [uM]', 'Discrete Nitrite [uM]',
                                               'Discrete Ammonium [uM]'),
                          'Discrete Nutrients Flag', 'Discrete Nutrients Replicate Flag', 'nutrients'),
    'salinity': Discrete('Salinity Samples', ('Discrete Salinity [psu]',), 'Discrete Salinity Flag',
                         'Discrete Salinity Replicate Flag', 'salinity'),
    'dic': Discrete('DIC/TA Samples', ('Discrete DIC [umol/kg]',), 'Discrete DIC Flag',
                    'Discrete DIC Replicate Flag', 'dic_ta'),
    'pco2': Discrete('DIC/TA Samples', ('Discrete pCO2 [uatm]', 'pCO2 Analysis Temp [deg C]'),
                     'Discrete pCO2 Flag', 'Discrete pCO2 Replicate Flag', 'dic_ta'),
    'alkalinity': Discrete(None, ('Discrete Alkalinity [umol/kg]',), 'Discrete Alkalinity Flag',
                           'Discrete Alkalinity Replicate Flag', None),
    'ph': Discrete(None, ('Discrete pH [Total scale]', 'pH Analysis Temp [deg C]'), 'Discrete pH Flag',
                   'Discrete pH Replicate Flag', None),
}
REPLICATE_BITS = {1: 2, 2: 3, 3: 4}

GENERAL_NOTES = [
    'Data sheets containing the discrete sampling estimates that we receive from the different laboratories that '
    'process the samples are provided "as is" for the user to review.',
    '',
    'We use the value -9999999 to represent missing values. The Discrete Data Flags (included in the spreadsheet '
    'and reproduced below) and the Summary Notes shown below can be used to determine why the missing value was set.']
SUMMARY_NOTES_INTRO = (
    'Notes are reported per Cruise ID, Cast Number, Bottle Number and Discrete Data Flag. Where note(s) are '
    'applicable to all casts and all bottles, the Cast Number and Bottle Number entries are set to "All Casts" and '
    '"All Bottles", respectively. Otherwise, specific numbers are used.')


class BuildError(ValueError):
    """The inputs cannot produce a summary."""


@dataclass
class BuildResult:
    """What a build produced."""
    rows: list[dict[str, str]]
    readme_notes: list[str]
    log: list[str] = field(default_factory=list)
    files: list[Path] = field(default_factory=list)


def template_columns() -> list[str]:
    with open(REFERENCE / 'discreteSummary_template.csv', newline='', encoding='utf-8-sig') as f:
        return next(csv.reader(f))


def fmt(value: str | float | None, dec: int) -> str:
    """Number as text with fixed decimals; FILL when missing."""
    return FILL if value in (None, '') else f'{float(value):.{dec}f}'


def sample_ids(cell) -> list[str]:
    """'49, 50, 051' -> ['49', '50', '51'] (numbers lose leading zeros)."""
    if cell in (None, ''):
        return []
    ids = [x.strip() for x in str(cell).split(',') if x.strip()]
    return [str(int(x)) if x.isdigit() else x for x in ids]


def load_labs(labs: dict[str, str | Path], cid: str, casts: dict, log: list[str]
              ) -> dict[str, dict[tuple[int | None, str], LabResult]]:
    """Read each lab file and index its results by (cast, sample number).

    The cast is None for labs whose files do not name it. Logs the file, its
    header facts, and station codes that disagree with the cast's site.

    Raises
    ------
    BuildError
        If the analysis is unknown or the file names a different cruise.
    """
    out: dict[str, dict[tuple[int | None, str], LabResult]] = {}
    for analysis, path in labs.items():
        if analysis not in DISCRETE:
            raise BuildError(f'unknown analysis {analysis!r}; expected one of {", ".join(DISCRETE)}')
        lab = read_lab(analysis, path)
        expedition = lab.header.get('Expedition')
        if expedition and expedition != cid:
            raise BuildError(f'{lab.path.name} is for {expedition}, not {cid}')
        facts = ', '.join(f'{k} {v}' for k, v in lab.header.items())
        log.append(f'{analysis}: {len(lab.results)} results from {lab.path.name}' + (f' ({facts})' if facts else ''))
        stations: dict[tuple[int, str], list[str]] = {}
        for res in lab.results:
            if res.station is not None and res.cast in casts:
                site = re.sub(r'\D', '', str(casts[res.cast]['Site']))
                if not site or int(res.station) != int(site):
                    stations.setdefault((res.cast, res.station), []).append(res.sample_id)
        for (c, code), ids in sorted(stations.items()):
            log.append(f'{analysis}: cast {c} is site {casts[c]["Site"]} but the lab station code is {code} '
                       f'(samples {", ".join(ids)}) -- WARNING, not used for the join')
        out[analysis] = {(res.cast, res.sample_id): res for res in lab.results}
    return out


def km_between(a: tuple[float, float], b: tuple[float, float]) -> float:
    la1, lo1, la2, lo2 = map(math.radians, (*a, *b))
    h = math.sin((la2 - la1) / 2) ** 2 + math.cos(la1) * math.cos(la2) * math.sin((lo2 - lo1) / 2) ** 2
    return 2 * 6371 * math.asin(math.sqrt(h))


def in_scope(spec, value: int) -> bool:
    """Notes sheet Cast/Niskin match: a number, a range such as '1-9', or All."""
    s = str(spec).strip()
    if s == 'All':
        return True
    if re.fullmatch(r'\d+\s*-\s*\d+', s):
        lo, hi = (int(x) for x in s.split('-'))
        return lo <= value <= hi
    return s == str(value)


def where(cast, niskin) -> str:
    """README wording: 'Cast 1, Bottle 9', 'Casts 1-9, All Bottles', ..."""
    c = 'All Casts' if str(cast) == 'All' else (f'Casts {cast}' if '-' in str(cast) else f'Cast {cast}')
    n = 'All Bottles' if str(niskin) == 'All' else f'Bottle {niskin}'
    return f'{c}, {n}'


def log_position(cast: dict) -> tuple[float, float]:
    """Cast Log start position in decimal degrees (west negative)."""
    lat = cast['Start Latitude [deg]'] + cast['Start Latitude [min]'] / 60
    deg, minutes = cast['Start Longitude [deg]'], cast['Start Longitude [min]']
    lon = deg - minutes / 60 if deg < 0 else -(deg + minutes / 60)   # EA is always west
    return lat, lon


def build(workbook: str | Path, btl_dir: str | Path, labs: dict[str, str | Path] | None = None) -> BuildResult:
    """Assemble the summary rows and published notes.

    Parameters
    ----------
    workbook : str or Path
        The cruise's transcription workbook (current layout).
    btl_dir : str or Path
        Folder holding one .btl per cast, named after the Cast Log CTD File.
    labs : dict, optional
        Analysis name (a DISCRETE key, e.g. 'oxygen') -> lab file. Analyses
        without a file stay "drawn, not received".

    Returns
    -------
    BuildResult

    Raises
    ------
    BuildError
        If a .btl is missing, a logged Niskin was not fired, a lab file names
        another cruise, a lab result puts a sample on a different Niskin than
        the workbook, or a lab result matches no drawn sample.
    """
    sheets = read_workbook(workbook)
    columns = template_columns()
    cruise = sheets['Cruise'][0]
    cid = cruise['Cruise']
    casts = {r['Cast']: r for r in sheets['Cast Log']}
    niskins = sorted(sheets['Niskin Log'], key=lambda r: (r['Cast'], r['Niskin']))
    issues = [r for r in sheets['Sample Issues'] if r['Issue'] == 'flag']
    notes = sheets['Notes']
    cast_bits, file_bits = bits_by_text('castFlags'), bits_by_text('CTDFileFlags')
    niskin_bits, sample_bits = bits_by_text('niskinFlags'), bits_by_text('discreteSampleFlags')
    result = BuildResult([], [])
    lab = load_labs(labs or {}, cid, casts, result.log)
    used: set[tuple[str, int | None, str]] = set()
    pending: dict[str, list[str]] = {}

    def note_bit(cast: int, niskin: int, column: str) -> list[int]:
        hit = any(n['Summary Flag Column'] == column and in_scope(n['Cast'], cast) and in_scope(n['Niskin'], niskin)
                  for n in notes)
        return [NOTES_BIT] if hit else []

    # ---- casts: .btl, start position/time, calibration age, published cast notes
    btl: dict[int, BtlFile] = {}
    start: dict[int, tuple[tuple[float, float], datetime]] = {}
    old: dict[int, list[str]] = {}
    for c, cast in casts.items():
        path = Path(btl_dir) / Path(str(cast['CTD File'])).with_suffix('.btl').name
        if not path.exists():
            raise BuildError(f'cast {c}: {path} not found')
        b = btl[c] = read_btl(path)
        t0 = datetime.combine(cast['Cast Date [UTC YYYY-MM-DD]'].date(), cast['Start Time [UTC HH:MM]'])
        if b.start:
            pos, d = b.start, km_between(b.start, log_position(cast))
            result.log.append(f'cast {c}: start position from .btl NMEA {pos[0]:.5f} {pos[1]:.5f}; Cast Log is '
                              f'{d:.2f} km away' + (' -- CHECK the transcription' if d > POSITION_TOLERANCE_KM else ''))
        else:
            pos = log_position(cast)
            result.log.append(f'cast {c}: no NMEA position in the .btl; start position taken from the Cast Log')
        if b.nmea_time:
            dt = abs((b.nmea_time - t0).total_seconds()) / 60
            if dt > TIME_TOLERANCE_MIN:
                result.log.append(f'cast {c}: Cast Log start time {t0:%H:%M} is {dt:.0f} min from the .btl NMEA time '
                                  f'{b.nmea_time:%H:%M:%S} -- CHECK')
        start[c] = (pos, t0)
        old[c] = sorted(s for s, d0 in b.calibrations.items() if (t0 - d0).days > CAL_MAX_AGE_DAYS)
        if old[c]:
            result.log.append(f'cast {c}: calibration more than a year before the cast: '
                              + ', '.join(f'{SENSOR_NAMES[s]} {b.calibrations[s]:%Y-%m-%d}' for s in old[c]))
        if cast['Notes'] and (cast['Cast Flag'] or cast['Cast Flag 2']):
            result.readme_notes.append(f'{cid}, {where(c, "All")}, Cast Flag -- {cast["Notes"]}')
        if cast['Notes'] and cast['CTD File Flag']:
            result.readme_notes.append(f'{cid}, {where(c, "All")}, CTD File Flag -- {cast["Notes"]}')

    # ---- one summary row per (cast, Niskin, replicate)
    logged = {(n['Cast'], n['Niskin']) for n in niskins}
    for n in niskins:
        c, k = n['Cast'], n['Niskin']
        cast, b = casts[c], btl[c]
        if k not in b.bottles:
            raise BuildError(f'cast {c} Niskin {k}: on the Niskin Log but not fired in {b.path.name}')
        bottle = b.bottles[k]
        (lat, lon), t0 = start[c]
        samples = {a: sample_ids(n[d.source]) if d.source else [] for a, d in DISCRETE.items()}
        nrep = max([1, *map(len, samples.values())])

        cbits = [cast_bits[f] for f in (cast['Cast Flag'], cast['Cast Flag 2']) if f] or [OK_BIT]
        if cast['Notes'] and (cast['Cast Flag'] or cast['Cast Flag 2']):
            cbits.append(NOTES_BIT)
        fbits = [file_bits[cast['CTD File Flag']]] if cast['CTD File Flag'] else [OK_BIT]
        if cast['Notes'] and cast['CTD File Flag']:
            fbits.append(NOTES_BIT)
        nbits = [niskin_bits[n['Niskin Flag']]] if n['Niskin Flag'] else [OK_BIT]
        if n['Niskin Flag'] and n['Notes']:
            nbits.append(NOTES_BIT)
            result.readme_notes.append(f'{cid}, {where(c, k)}, Niskin Flag -- {n["Notes"]}')

        for i in range(nrep):
            r = dict.fromkeys(columns, FILL)
            r.update({
                'Cruise': cid, 'Station': cast['Site'], 'Target Asset': cast['Target Assets'] or FILL,
                'Start Latitude [degrees]': f'{lat:.5f}', 'Start Longitude [degrees]': f'{lon:.5f}',
                'Start Time [UTC]': t0.strftime('%Y-%m-%dT%H:%M:%S.000Z'), 'Cast': str(c),
                'Cast Flag': flag(*cbits, *note_bit(c, k, 'Cast Flag')),
                'Bottom Depth at Start Position [m]': fmt(cast['Start Water Depth [m]'], 1).removesuffix('.0'),
                'CTD File': cast['CTD File'],
                'CTD File Flag': flag(*fbits, *note_bit(c, k, 'CTD File Flag')),
                'Niskin/Bottle Position': str(k),
                'Niskin Flag': flag(*nbits, *note_bit(c, k, 'Niskin Flag')),
                'CTD Bottle Closure Time [UTC]': bottle['time'].strftime('%Y-%m-%dT%H:%M:%S.000Z'),
            })
            for var, (col, dec) in CTD_COLUMNS.items():
                r[col] = fmt(bottle.get(var), dec)
            for col, (sensor, analysis) in CTD_FLAGS.items():
                if col == 'CTD pH Flag' and 'Ph' not in b.names:
                    continue
                bits = [OK_BIT, *note_bit(c, k, col)]
                if sensor in old[c]:
                    bits.append(OLD_CAL_BIT)
                if analysis and i < len(samples[analysis]):
                    bits.append(MATCH_BIT)
                r[col] = flag(*bits)
            for a, d in DISCRETE.items():
                got = samples[a]
                if i < len(got):
                    issue_bits = []
                    for s in issues:
                        if ((s['Cast'], s['Niskin']) == (c, k) and s['Analysis'] == d.issue_name
                                and sample_ids(s['Sample ID']) == [got[i]]):
                            issue_bits += [sample_bits[s['Sample Flag']], NOTES_BIT]
                    table = lab.get(a, {})
                    key = (c, got[i]) if (c, got[i]) in table else (None, got[i])
                    res = table.get(key)
                    if res is None:
                        bits = [NOT_RECEIVED_BIT, *issue_bits]
                        if a in lab:
                            pending.setdefault(a, []).append(f'cast {c} Niskin {k} sample {got[i]}')
                    else:
                        if res.niskin is not None and res.niskin != k:
                            raise BuildError(f'{a} sample {got[i]}: lab file line {res.line} puts it on cast {c} '
                                             f'Niskin {res.niskin}, the workbook on Niskin {k}')
                        used.add((a, *key))
                        r.update(res.values)
                        bits = issue_bits or [OK_BIT]
                        logged_temp = n.get('Oxygen Draw Temperature [deg C]') if a == 'oxygen' else None
                        if (res.draw_temp is not None and logged_temp not in (None, '')
                                and abs(res.draw_temp - float(logged_temp)) > DRAW_TEMP_TOLERANCE_C):
                            result.log.append(f'{a}: cast {c} Niskin {k} sample {got[i]}: lab draw temperature '
                                              f'{res.draw_temp} vs log {logged_temp} -- CHECK (log kept)')
                        if res.comment:
                            result.log.append(f'{a}: cast {c} Niskin {k} sample {got[i]}: lab comment '
                                              f'"{res.comment}" -- not flagged; meaning not yet confirmed')
                    r[d.flag] = flag(*bits, *note_bit(c, k, d.flag))
                    r[d.replicate] = flag(REPLICATE_BITS[len(got)])
                else:
                    r[d.flag] = flag(NOT_DRAWN_BIT, *note_bit(c, k, d.flag))
            result.rows.append(r)

    # ---- lab results must all land on a drawn sample; drawn samples without one are logged
    for a, table in lab.items():
        extra = [res for key, res in table.items() if (a, *key) not in used]
        if extra:
            raise BuildError(f'{a}: lab results match no sample on the Niskin Log: ' + '; '.join(
                f'line {res.line} cast {res.cast} Niskin {res.niskin} sample {res.sample_id}' for res in extra))
        if a in pending:
            result.log.append(f'{a}: drawn but no lab result (flag bit 1): {"; ".join(pending[a])}')

    # ---- published notes: Sample Issues flag rows, Notes sheet, calibration age
    for s in issues:
        if (s['Cast'], s['Niskin']) not in logged:
            raise BuildError(f'Sample Issues: cast {s["Cast"]} Niskin {s["Niskin"]} is not on the Niskin Log')
        col = next(d.flag for d in DISCRETE.values() if d.issue_name == s['Analysis'])
        result.readme_notes.append(f'{cid}, {where(s["Cast"], s["Niskin"])}, {col} -- {s["Reason"]} '
                                   f'(sample {s["Sample ID"]})')
    for n in notes:
        result.readme_notes.append(f'{cid}, {where(n["Cast"], n["Niskin"])}, {n["Summary Flag Column"]} -- '
                                   f'{n["Note"]}')
    old_all = sorted({s for v in old.values() for s in v})
    if old_all:
        cols = [col for col, (sensor, _) in CTD_FLAGS.items() if sensor in old_all]
        first = {s: next(b.calibrations[s] for b in btl.values() if s in b.calibrations) for s in old_all}
        dates = ', '.join(f'{SENSOR_NAMES[s]} {d0:%Y-%m-%d}' for s, d0 in first.items())
        result.readme_notes.append(f'{cid}, All Casts, All Bottles, {", ".join(cols)} -- Calibration coefficients were '
                                   f'more than one year old at the time of the cruise ({dates}); bit 6 is set.')
    result.log.insert(0, f'{len(result.rows)} rows from {len(niskins)} Niskins on {len(casts)} casts')
    return result


def write_xlsx(path: Path, columns: list[str], rows: list[dict[str, str]]) -> None:
    """Formatted summary workbook plus the flag definitions sheet."""
    plain, bold = Font(name='Arial', size=10), Font(name='Arial', size=10, bold=True)
    centre = Alignment(horizontal='center', vertical='center')
    rule = Border(bottom=Side(style='thin', color='000000'))
    wb = Workbook()
    ws = wb.active
    ws.title = 'Discrete Summary'
    text_cols = {'Cruise', 'Station', 'Target Asset', 'CTD File'}
    for j, col in enumerate(columns, start=1):
        c = ws.cell(row=1, column=j, value=col)
        c.font, c.border = bold, rule
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        width = 44 if col == 'Target Asset' else 24 if '[UTC]' in col else 19.7 if col.endswith('Flag') else 14
        ws.column_dimensions[c.column_letter].width = width
    for r, row in enumerate(rows, start=2):
        for j, col in enumerate(columns, start=1):
            text = row[col]
            c = ws.cell(row=r, column=j)
            if col not in text_cols and re.fullmatch(r'-?\d+(\.\d+)?', text):
                dec = len(text.split('.')[1]) if '.' in text else 0
                c.value = float(text) if dec else int(text)
                c.number_format = '0.' + '0' * dec if dec else '0'
            else:
                c.value, c.number_format = text, '@'
            c.font, c.alignment = plain, centre
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = 'A2'

    fs = wb.create_sheet('Discrete Data Flags')
    for k, (key, title) in enumerate(FLAG_TYPES.items()):
        for j, h in enumerate(('Bit Position', title)):
            c = fs.cell(row=1, column=2 * k + j + 1, value=h)
            c.font, c.border = bold, rule
        fs.column_dimensions[c.column_letter].width = 48
        for r, f in enumerate(definitions(), start=2):
            fs.cell(row=r, column=2 * k + 1, value=int(f['bitPosition'])).font = plain
            fs.cell(row=r, column=2 * k + 2, value=f[key]).font = plain
    fs.freeze_panes = 'A2'
    wb.save(path)


def write_readme(path: Path, title: str, history: list[str], notes: list[str]) -> None:
    """Summary README in the historic EA layout (LF line endings)."""
    out = [title, '', 'Revision History:', '', *history, '', 'General Notes:', '', *GENERAL_NOTES, '',
           'Summary Notes:', '', SUMMARY_NOTES_INTRO, '']
    for n in notes:
        out += [n, '']
    out += ['Discrete Data Flags:', '']
    for key, name in FLAG_TYPES.items():
        out.append(f'* {name}:\t\tBit \tFlags')
        out += [f'\t\t\t\t\t{f["bitPosition"]}\t\t{f[key]}' for f in definitions()]
        out.append('')
    path.write_text('\n'.join(out) + '\n', encoding='utf-8', newline='\n')


def read_history(readme: Path) -> list[str]:
    """Revision History lines of an existing README (empty if none)."""
    if not readme.exists():
        return []
    lines = readme.read_text(encoding='utf-8').splitlines()
    try:
        i = lines.index('Revision History:')
    except ValueError:
        return []
    out = []
    for ln in lines[i + 1:]:
        if ln.startswith('General Notes:'):
            break
        if ln.strip():
            out.append(ln)
    return out


def write_outputs(result: BuildResult, workbook: str | Path, out_dir: str | Path, version: str,
                  note: str) -> list[Path]:
    """Write the xlsx, README (appending to its revision history) and build log.

    Raises
    ------
    BuildError
        If `version` is already in the README's revision history.
    """
    cruise = read_workbook(workbook)['Cruise'][0]
    expedition = str(cruise['Expedition'])
    stem = f'{expedition.replace(" ", "-")}_{cruise["Cruise"]}_Discrete_Summary'
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    readme = out / f'{stem}_README.txt'
    history = read_history(readme)
    if any(h.split(' ', 1)[0] == version for h in history):
        raise BuildError(f'version {version} is already in {readme.name}; use a new version number')
    history.append(f'{version} ({date.today():%Y-%m-%d}) -- {note}')
    title = (f'{expedition.replace("-", " ")} ({cruise["Ship"]}, {cruise["Start Date [UTC YYYY-MM-DD]"]:%Y-%m-%d} '
             f'to {cruise["End Date [UTC YYYY-MM-DD]"]:%Y-%m-%d}) Discrete Sampling README')
    xlsx = out / f'{stem}.xlsx'
    write_xlsx(xlsx, template_columns(), result.rows)
    write_readme(readme, title, history, result.readme_notes)
    log = out / f'{stem}_build_log.txt'
    log.write_text('\n'.join([f'{stem} {version} built {datetime.now():%Y-%m-%d %H:%M}', *result.log]) + '\n',
                   encoding='utf-8', newline='\n')
    result.files = [xlsx, readme, log]
    return result.files
