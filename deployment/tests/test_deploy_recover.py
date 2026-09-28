"""End-to-end tests for the deploy and recover subcommands."""
from __future__ import annotations

import csv
import os
import shutil
from pathlib import Path

import openpyxl
import pytest

from asset_deploy import cli, cruise
from asset_deploy.config_xlsx import ConfigError, ConfigTable
from asset_deploy.metadata import read_metadata

FIXTURES = Path(__file__).parent / 'fixtures'
_REAL_AM = Path(os.environ.get('ASSET_MANAGEMENT_DIR', r'C:\\Users\\chris\\code\\asset-management'))

DEPLOY_HEADER = [
    'deploymentNumber', 'Reference Designator', 'sensor.uid', 'startDateTime', 'stopDateTime', 'lat', 'lon',
    'mooring.uid', 'CUID_Deploy', 'node.uid', 'versionNumber', 'deployedBy', 'CUID_Recover', 'recoveredBy',
    'orbit', 'deployment_depth', 'water_depth', 'notes', 'electrical.uid', 'assembly_template_revision',
]
CRUISES = [
    ['CUID', 'ShipName', 'cruiseStartDateTime', 'cruiseStopDateTime', 'notes'],
    ['UNRECOVERED', 'N/A', '2009-09-01T00:00:00', '2040-09-30T00:00:00', 'placeholder'],
    ['SR2614', 'Sally Ride', '2026-09-30T00:00:00', '2026-10-02T23:59:59', 'Endurance 25'],
    ['SR2703', 'Sally Ride', '2027-04-10T00:00:00', '2027-04-14T23:59:59', 'Endurance 26'],
]


def _write_csv(path: Path, rows: list[list[str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, 'w', newline='', encoding='utf-8') as f:
        csv.writer(f, lineterminator='\n').writerows(rows)


def _read_rows(path: Path) -> list[dict[str, str]]:
    with open(path, newline='', encoding='utf-8') as f:
        return list(csv.DictReader(f))


@pytest.fixture
def small_am(tmp_path: Path) -> Path:
    """asset-management with one CE07SHSM deployment (21) of two rows."""
    _write_csv(tmp_path / 'cruise' / 'CruiseInformation.csv', CRUISES)
    base = dict.fromkeys(DEPLOY_HEADER, '')
    base.update({'deploymentNumber': '21', 'startDateTime': '2026-10-01T12:00:00', 'versionNumber': '1'})
    deployed = {**base, 'Reference Designator': 'CE07SHSM-RID27-03-CTDBPC000', 'CUID_Deploy': 'SR2614'}
    skipped = {**base, 'Reference Designator': 'CE07SHSM-MFD37-01-OPTAAD000', 'CUID_Deploy': '#SR2614'}
    _write_csv(tmp_path / 'deployment' / 'CE07SHSM_Deploy.csv',
               [DEPLOY_HEADER] + [[r[k] for k in DEPLOY_HEADER] for r in (deployed, skipped)])
    return tmp_path


def _recover(am: Path, *extra: str) -> None:
    cli.main(['recover', 'CE07SHSM-00021', '--cruise', 'SR2703', '--time', '2027-04-11T15:30',
              '--asset-management', str(am), *extra])


# -- platform from the sheet ---------------------------------------------------------------------------------------

def _sheet(refdes: str, mooring_sn: str) -> ConfigTable:
    return ConfigTable([
        ['Reference Designator:', refdes, ''],
        ['Mooring S/N:', mooring_sn, ''],
        ['Anchor Launch Date:', '2026-10-01', ''],
        ['Anchor Launch Time:', '12:00:00', ''],
        ['Latitude:', '46.0', ''],
        ['Longitude:', '-124.5', ''],
        ['Water Depth:', '87 m', ''],
        ['Cruise #:', 'SR2614', ''],
    ])


def test_platform_read_from_sheet():
    assert read_metadata(_sheet('ce07shsm', 'CE07SHSM-00021')).platform_code == 'CE07SHSM'


def test_sheet_rows_disagree():
    with pytest.raises(ConfigError, match='do not agree'):
        read_metadata(_sheet('CE07SHSM', 'CE02SHSM-00021'))


# -- cruise check --------------------------------------------------------------------------------------------------

def test_cruise_known_and_in_window(small_am: Path):
    assert cruise.check_cruise(small_am, 'SR2614', '2026-10-01T12:00:00', 'Launch').ship == 'Sally Ride'
    assert cruise.check_cruise(small_am, '#SR2614', '2026-10-01T12:00', 'Launch').cuid == 'SR2614'
    assert cruise.check_cruise(small_am, 'UNRECOVERED', '2031-01-01T00:00:00', 'Recovery').cuid == 'UNRECOVERED'


def test_cruise_unknown(small_am: Path):
    with pytest.raises(ValueError, match="'SR2699' is not in cruise/CruiseInformation.csv"):
        cruise.check_cruise(small_am, 'SR2699', '2026-10-01T12:00:00', 'Launch')


def test_cruise_time_outside_window(small_am: Path):
    with pytest.raises(ValueError, match='Launch time 2026-10-05T12:00:00 is outside cruise SR2614'):
        cruise.check_cruise(small_am, 'SR2614', '2026-10-05T12:00:00', 'Launch')


# -- recover -------------------------------------------------------------------------------------------------------

def test_recover_writes_every_row_and_keeps_not_deployed_marker(small_am: Path):
    _recover(small_am)
    rows = _read_rows(small_am / 'deployment' / 'CE07SHSM_Deploy.csv')
    assert [r['stopDateTime'] for r in rows] == ['2027-04-11T15:30:00'] * 2
    assert [r['CUID_Recover'] for r in rows] == ['SR2703', '#SR2703']


def test_recover_refuses_to_overwrite_without_force(small_am: Path):
    _recover(small_am)
    with pytest.raises(SystemExit, match='already has recovery info on 2 of 2 rows'):
        cli.main(['recover', 'CE07SHSM-00021', '--cruise', 'SR2703', '--time', '2027-04-12T09:00:00',
                  '--asset-management', str(small_am)])
    cli.main(['recover', 'CE07SHSM-00021', '--cruise', 'SR2703', '--time', '2027-04-12T09:00:00',
              '--asset-management', str(small_am), '--force'])
    rows = _read_rows(small_am / 'deployment' / 'CE07SHSM_Deploy.csv')
    assert {r['stopDateTime'] for r in rows} == {'2027-04-12T09:00:00'}


def test_recover_before_start_rejected(small_am: Path):
    with pytest.raises(SystemExit, match='not after deployment start 2026-10-01T12:00:00'):
        cli.main(['recover', 'CE07SHSM-00021', '--cruise', 'SR2614', '--time', '2026-10-01T06:00:00',
                  '--asset-management', str(small_am)])


def test_recover_unknown_deployment(small_am: Path):
    with pytest.raises(SystemExit, match='no rows for deployment 22'):
        cli.main(['recover', 'CE07SHSM-00022', '--cruise', 'SR2703', '--time', '2027-04-11T15:30:00',
                  '--asset-management', str(small_am)])


def test_recover_test_mode_leaves_original(small_am: Path):
    _recover(small_am, '--mode', 'test')
    assert {r['stopDateTime'] for r in _read_rows(small_am / 'deployment' / 'CE07SHSM_Deploy.csv')} == {''}
    preview = small_am / 'deployment' / '_preview_CE07SHSM_Deploy.csv'
    assert {r['stopDateTime'] for r in _read_rows(preview)} == {'2027-04-11T15:30:00'}


# -- deploy (needs the real bulk asset records) --------------------------------------------------------------------

@pytest.fixture
def real_am(tmp_path: Path) -> Path:
    """Copy of the real asset-management records minus CE02SHSM-00021."""
    if not (_REAL_AM / 'bulk').exists():
        pytest.skip('asset-management repo not mounted')
    for name in ('sensor', 'eng', 'platform'):
        rel = Path('bulk') / f'{name}_bulk_load-AssetRecord.csv'
        (tmp_path / 'bulk').mkdir(exist_ok=True)
        lines = (_REAL_AM / rel).read_text(encoding='utf-8').splitlines()
        (tmp_path / rel).write_text(
            '\n'.join(ln for ln in lines if not ln.startswith('CGMCE-02SHSM-00021,')) + '\n', encoding='utf-8')
    _write_csv(tmp_path / 'cruise' / 'CruiseInformation.csv', CRUISES)
    rows = _read_rows(_REAL_AM / 'deployment' / 'CE02SHSM_Deploy.csv')
    header = list(rows[0])
    _write_csv(tmp_path / 'deployment' / 'CE02SHSM_Deploy.csv',
               [header] + [[r[k] for k in header] for r in rows if r['deploymentNumber'] != '21'])
    shutil.copy(FIXTURES / 'CE02SHSM-00021.xlsx', tmp_path / 'CE02SHSM-00021.xlsx')
    return tmp_path


def test_deploy_end_to_end(real_am: Path):
    cli.main(['deploy', '--config', str(real_am / 'CE02SHSM-00021.xlsx'), '--asset-management', str(real_am)])
    new = [r for r in _read_rows(real_am / 'deployment' / 'CE02SHSM_Deploy.csv') if r['deploymentNumber'] == '21']
    assert len(new) == 30
    assert {r['CUID_Deploy'] for r in new} == {'SR2614'}
    uids = {r['sensor.uid'] for r in new}
    assert {'CGINS-OPTAAJ-00284', 'CGINS-PHSENG-02102'} <= uids
    platform = (real_am / 'bulk' / 'platform_bulk_load-AssetRecord.csv').read_text(encoding='utf-8')
    assert '\nCGMCE-02SHSM-00021,' in platform

    with pytest.raises(SystemExit, match='already has rows for deployment 21; use --replace'):
        cli.main(['deploy', '--config', str(real_am / 'CE02SHSM-00021.xlsx'), '--asset-management', str(real_am)])


def test_deploy_unknown_cruise(real_am: Path):
    _write_csv(real_am / 'cruise' / 'CruiseInformation.csv', CRUISES[:2])
    with pytest.raises(SystemExit, match="'SR2614' is not in"):
        cli.main(['deploy', '--config', str(real_am / 'CE02SHSM-00021.xlsx'), '--asset-management', str(real_am)])
    assert not [r for r in _read_rows(real_am / 'deployment' / 'CE02SHSM_Deploy.csv')
                if r['deploymentNumber'] == '21']


def _edit_sheet(path: Path, edits: dict[str, str]) -> None:
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    for row in ws.iter_rows(min_col=1, max_col=2):
        if row[0].value in edits:
            row[1].value = edits.pop(row[0].value)
    assert not edits, f'labels not found: {edits}'
    wb.save(path)


def _deploy(am: Path, *extra: str) -> None:
    cli.main(['deploy', '--config', str(am / 'CE02SHSM-00021.xlsx'), '--asset-management', str(am), *extra])


def _dep21(am: Path) -> dict[str, dict[str, str]]:
    rows = _read_rows(am / 'deployment' / 'CE02SHSM_Deploy.csv')
    return {r['Reference Designator']: r for r in rows if r['deploymentNumber'] == '21'}


def test_replace_updates_placeholders_swaps_and_failures(real_am: Path, capsys: pytest.CaptureFixture):
    _deploy(real_am)
    before = (real_am / 'deployment' / 'CE02SHSM_Deploy.csv').read_text(encoding='utf-8').splitlines()
    _edit_sheet(real_am / 'CE02SHSM-00021.xlsx', {
        'Latitude:': '44.63612', 'Longitude:': '-124.30421', 'Anchor Launch Time:': '17:42:00',
        'PHSEN S/N': '721-2087',        # swapped on deck
        'SPKIR S/N': 'Not Deployed',    # failed before deployment
    })
    capsys.readouterr()
    _deploy(real_am, '--replace')
    out = capsys.readouterr().out
    assert "lat: '45.50000' -> '44.63612'  (all rows)" in out
    assert "sensor.uid: 'CGINS-PHSENG-02102' -> 'CGINS-PHSENG-02087'  (RID26-06-PHSENG000)" in out

    rows = _dep21(real_am)
    assert len(rows) == 30
    assert {r['startDateTime'] for r in rows.values()} == {'2026-09-30T17:42:00'}
    spkir = rows['CE02SHSM-RID26-08-SPKIRB000']
    assert (spkir['sensor.uid'], spkir['CUID_Deploy'], spkir['notes']) == ('', '#SR2614', 'Not Deployed')

    after = (real_am / 'deployment' / 'CE02SHSM_Deploy.csv').read_text(encoding='utf-8').splitlines()
    assert len(after) == len(before)
    assert [a for a, b in zip(after, before) if not a.startswith('21,')] == \
           [b for b in before if not b.startswith('21,')]


def test_replace_without_changes(real_am: Path, capsys: pytest.CaptureFixture):
    _deploy(real_am)
    capsys.readouterr()
    _deploy(real_am, '--replace')
    assert 'No changes from the existing rows.' in capsys.readouterr().out


def test_replace_test_mode_leaves_original(real_am: Path):
    _deploy(real_am)
    _edit_sheet(real_am / 'CE02SHSM-00021.xlsx', {'Latitude:': '44.63612'})
    _deploy(real_am, '--replace', '--mode', 'test')
    assert {r['lat'] for r in _dep21(real_am).values()} == {'45.50000'}
    preview = _read_rows(real_am / 'deployment' / '_preview_CE02SHSM_Deploy.csv')
    assert {r['lat'] for r in preview if r['deploymentNumber'] == '21'} == {'44.63612'}


def test_replace_requires_existing_deployment(real_am: Path):
    with pytest.raises(SystemExit, match='no rows for deployment 21 to replace'):
        _deploy(real_am, '--replace')


def test_replace_refuses_recovered_deployment(real_am: Path):
    _deploy(real_am)
    cli.main(['recover', 'CE02SHSM-00021', '--cruise', 'SR2614', '--time', '2026-10-02T12:00:00',
              '--asset-management', str(real_am)])
    with pytest.raises(SystemExit, match='already recovered'):
        _deploy(real_am, '--replace')
